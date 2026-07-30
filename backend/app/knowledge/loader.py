from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Callable

from app.knowledge.document import Document


class DocumentLoader:
    """
    企业文档加载器。

    支持 PDF / Markdown / TXT / DOCX，输出统一 Document（content + metadata）。
    """

    _TEXT_EXTENSIONS = frozenset({".txt", ".md", ".markdown"})
    _PDF_EXTENSIONS = frozenset({".pdf"})
    _DOCX_EXTENSIONS = frozenset({".docx"})

    def load(
        self,
        path: str | Path,
        *,
        extra_metadata: dict | None = None,
    ) -> Document:
        """
        从文件路径加载单个文档。
        """

        file_path = Path(path).expanduser().resolve()

        if not file_path.is_file():

            raise FileNotFoundError(
                f"Document file not found: {file_path}"
            )

        suffix = file_path.suffix.lower()

        if suffix in self._TEXT_EXTENSIONS:

            content = self._load_text(file_path)

        elif suffix in self._PDF_EXTENSIONS:

            content = self._load_pdf(file_path)

        elif suffix in self._DOCX_EXTENSIONS:

            content = self._load_docx(file_path)

        else:

            raise ValueError(
                f"Unsupported file type '{suffix}' for {file_path.name}. "
                f"Supported: txt, md, markdown, pdf, docx"
            )

        document_id = hashlib.sha256(
            str(file_path).encode("utf-8")
        ).hexdigest()[:16]

        metadata = {
            "document_id": document_id,
            "source_path": str(file_path),
            "file_name": file_path.name,
            "file_type": suffix.lstrip("."),
            "loader": "DocumentLoader",
        }

        if extra_metadata:

            metadata.update(extra_metadata)

        return Document(
            content=content.strip(),
            metadata=metadata,
        )

    def load_many(
        self,
        paths: list[str | Path],
    ) -> list[Document]:

        return [self.load(path) for path in paths]

    def load_directory(
        self,
        directory: str | Path,
        *,
        recursive: bool = False,
    ) -> list[Document]:
        """
        加载目录下支持的文件。
        """

        root = Path(directory).expanduser().resolve()

        if not root.is_dir():

            raise NotADirectoryError(
                f"Not a directory: {root}"
            )

        pattern = "**/*" if recursive else "*"

        documents: list[Document] = []

        for file_path in sorted(root.glob(pattern)):

            if not file_path.is_file():

                continue

            suffix = file_path.suffix.lower()

            if (
                suffix
                in self._TEXT_EXTENSIONS
                | self._PDF_EXTENSIONS
                | self._DOCX_EXTENSIONS
            ):

                documents.append(
                    self.load(file_path)
                )

        return documents

    def _load_text(
        self,
        file_path: Path,
    ) -> str:

        return file_path.read_text(encoding="utf-8")

    def _load_pdf(
        self,
        file_path: Path,
    ) -> str:

        reader_factory = self._pdf_reader_factory()

        reader = reader_factory(file_path)

        pages: list[str] = []

        for page in reader.pages:

            text = page.extract_text()

            if text:

                pages.append(text)

        return "\n\n".join(pages)

    def _load_docx(
        self,
        file_path: Path,
    ) -> str:

        doc = self._docx_document_factory(file_path)

        paragraphs = [
            paragraph.text.strip()
            for paragraph in doc.paragraphs
            if paragraph.text.strip()
        ]

        return "\n\n".join(paragraphs)

    @staticmethod
    def _pdf_reader_factory() -> Callable:
        """
        延迟导入 pypdf，未安装时给出安装提示。
        """

        try:

            from pypdf import PdfReader

        except ImportError as error:

            raise RuntimeError(
                "PDF loading requires pypdf. "
                "Install with: pip install pypdf"
            ) from error

        return PdfReader

    @staticmethod
    def _docx_document_factory(
        file_path: Path,
    ):
        """
        延迟导入 python-docx。
        """

        try:

            from docx import Document as DocxDocument

        except ImportError as error:

            raise RuntimeError(
                "DOCX loading requires python-docx. "
                "Install with: pip install python-docx"
            ) from error

        return DocxDocument(str(file_path))
