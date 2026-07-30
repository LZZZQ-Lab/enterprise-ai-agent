"""企业知识库 Document 加载与切分测试。"""

from __future__ import annotations

from pathlib import Path

import pytest

from app.embedding.fake import FakeEmbedding
from app.knowledge.document import Document
from app.knowledge.loader import DocumentLoader
from app.knowledge.splitter import DocumentSplitter
from app.rag.knowledge_base import KnowledgeBase
from app.rag.vectorstore.inmemory import InMemoryVectorStore

from tests.fixtures.paths import KNOWLEDGE_FIXTURES

FIXTURES = KNOWLEDGE_FIXTURES


def test_document_to_dict_shape() -> None:

    doc = Document(
        content="hello",
        metadata={"document_id": "t1", "source_path": "/x"},
    )

    payload = doc.to_dict()

    assert set(payload.keys()) == {"content", "metadata"}
    assert payload["content"] == "hello"
    assert payload["metadata"]["document_id"] == "t1"


def test_load_txt_and_md() -> None:

    loader = DocumentLoader()

    txt = loader.load(FIXTURES / "sample.txt")
    md = loader.load(FIXTURES / "sample.md")

    assert "Enterprise AI Agent Platform" in txt.content
    assert txt.metadata["file_type"] == "txt"
    assert txt.metadata["source_path"]

    assert "Agent Runtime" in md.content
    assert md.metadata["file_type"] == "md"


def test_splitter_produces_chunks() -> None:

    loader = DocumentLoader()
    doc = loader.load(FIXTURES / "sample.md")

    splitter = DocumentSplitter()
    chunks = splitter.split(doc)

    assert len(chunks) >= 2

    for index, chunk in enumerate(chunks):

        assert chunk.metadata["chunk_index"] == index
        assert chunk.metadata["parent_id"] == doc.document_id
        assert chunk.content


def test_load_directory() -> None:

    loader = DocumentLoader()
    docs = loader.load_directory(FIXTURES)

    names = {doc.metadata["file_name"] for doc in docs}

    assert "sample.txt" in names
    assert "sample.md" in names


def test_pipeline_loader_splitter_to_rag_ingest() -> None:

    loader = DocumentLoader()
    splitter = DocumentSplitter()

    doc = loader.load(FIXTURES / "sample.txt")
    chunks = splitter.split(doc)

    embedding = FakeEmbedding(dimension=32)
    store = InMemoryVectorStore(dimension=32)
    kb = KnowledgeBase(
        embedding_provider=embedding,
        vector_store=store,
    )

    for chunk in chunks:

        kb.ingest(chunk.to_rag_document())

    hits = kb.search("RAG 知识库", top_k=2)

    assert hits


@pytest.fixture
def sample_pdf_path(tmp_path: Path) -> Path:
    """运行时生成最小 PDF 测试文件。"""

    pytest.importorskip("pypdf")

    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)

    path = tmp_path / "sample.pdf"
    writer.write(path)

    return path


def test_load_pdf_file(sample_pdf_path: Path) -> None:

    pytest.importorskip("pypdf")

    loader = DocumentLoader()
    doc = loader.load(sample_pdf_path)

    assert doc.metadata["file_type"] == "pdf"
    assert isinstance(doc.content, str)


@pytest.fixture
def sample_docx_path(tmp_path: Path) -> Path:
    """运行时生成 DOCX 测试文件。"""

    docx = pytest.importorskip("docx")

    from docx import Document as DocxDocument

    document = DocxDocument()
    document.add_paragraph("企业 DOCX 示例段落。")
    document.add_paragraph("第二段：知识库导入测试。")

    path = tmp_path / "sample.docx"
    document.save(path)

    return path


def test_load_docx_file(sample_docx_path: Path) -> None:

    pytest.importorskip("docx")

    loader = DocumentLoader()
    doc = loader.load(sample_docx_path)

    assert "企业 DOCX" in doc.content
    assert "知识库导入" in doc.content
    assert doc.metadata["file_type"] == "docx"


def test_unsupported_extension(tmp_path: Path) -> None:

    bad = tmp_path / "data.xyz"
    bad.write_text("x", encoding="utf-8")

    loader = DocumentLoader()

    with pytest.raises(ValueError, match="Unsupported file type"):

        loader.load(bad)
